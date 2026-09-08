"""DOCX extraction.

python-docx gives us the document tree, so unlike PDF we can see tables,
headers and footers directly rather than inferring them. Table cells and
header/footer text are the two things ATS parsers most often drop, so we both
record them as layout warnings and still include their text (a candidate's
skills matrix is often inside a table and should count toward matching).
"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.extraction.base import ExtractedDocument, PageStats, SourceFormat
from app.errors import ExtractionFailed

# Word does not store page count reliably; ~500 words per page is the standard
# approximation used for resume length guidance.
WORDS_PER_PAGE = 500


def extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:
        raise ExtractionFailed(
            "The DOCX could not be parsed. If this is a legacy .doc file, "
            "re-save it as .docx or PDF."
        ) from exc

    warnings: list[str] = []
    blocks: list[str] = []
    table_count = 0

    for element in _iter_block_items(document):
        if isinstance(element, Paragraph):
            text = element.text.strip()
            if text:
                blocks.append(_prefix_for(element, text))
        else:
            table_count += 1
            blocks.extend(_flatten_table(element))

    header_footer_text = _header_footer_text(document)
    if header_footer_text:
        warnings.append(
            "Contact details or content sit in the document header/footer; "
            "many ATS parsers discard those regions entirely."
        )
        blocks.extend(header_footer_text)

    if table_count:
        warnings.append(
            f"Document uses {table_count} table(s); ATS parsers often read table "
            "cells out of order or drop them."
        )

    image_count = len(document.inline_shapes)
    if image_count:
        warnings.append(
            f"Document embeds {image_count} image(s); text inside images is invisible to an ATS."
        )

    text = "\n".join(blocks).strip()
    if not text:
        raise ExtractionFailed("The DOCX contains no readable text.")

    word_count = len(text.split())
    page_count = max(1, round(word_count / WORDS_PER_PAGE) or 1)

    return ExtractedDocument(
        text=text,
        source_format=SourceFormat.DOCX,
        page_count=page_count,
        pages=[
            PageStats(index=0, char_count=len(text), line_count=len(blocks), has_extractable_text=True)
        ],
        warnings=warnings,
        has_tables=table_count > 0,
        table_count=table_count,
        has_images=image_count > 0,
        image_count=image_count,
        has_text_in_headers_footers=bool(header_footer_text),
    )


def _prefix_for(paragraph: Paragraph, text: str) -> str:
    """Re-attach a bullet marker to list paragraphs.

    Word stores list bullets as numbering properties, not characters. Without
    this, every bullet arrives as an ordinary sentence and the content-quality
    scorer cannot tell bullets from prose.
    """
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    is_list = "list" in style
    if not is_list:
        p_pr = paragraph._p.pPr  # noqa: SLF001 - python-docx exposes no public accessor
        is_list = p_pr is not None and p_pr.find(qn("w:numPr")) is not None
    if is_list and not text.startswith(("-", "*")):
        return f"- {text}"
    return text


def _flatten_table(table: Table) -> list[str]:
    """Read a table row-major, which is how a human reads a skills matrix."""
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        # De-duplicate horizontally merged cells, which repeat their text.
        deduped: list[str] = []
        for cell in cells:
            if cell and (not deduped or deduped[-1] != cell):
                deduped.append(cell)
        if deduped:
            rows.append(" | ".join(deduped))
    return rows


def _header_footer_text(document: DocxDocument) -> list[str]:
    found: list[str] = []
    for section in document.sections:
        for part in (section.header, section.footer):
            if part is None:
                continue
            for paragraph in part.paragraphs:
                text = paragraph.text.strip()
                if text:
                    found.append(text)
    return found


def _iter_block_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in true document order."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)
